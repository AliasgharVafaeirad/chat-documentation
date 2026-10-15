import tempfile
from pathlib import Path
import unittest
import json
from unittest.mock import patch
from openpyxl import load_workbook
import process_backup as app


class PipelineTests(unittest.TestCase):
    def test_empty_item_does_not_discard_good_item(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            (folder / "downloads").mkdir()
            messages = []
            raw = {}
            for i in range(2):
                ref = f"downloads/{i}.ogg"
                (folder / ref).write_bytes(str(i).encode())
                messages.append({"id": str(i), "text": "", "attachment_results": [ref]})
                raw[ref] = {
                    "status": "done",
                    "sha256": app.digest(folder / ref),
                    "text": "voice text",
                }
            app.save(folder / "messages.json", messages)
            app.save(folder / "transcription_results.json", raw)

            def request(command, target):
                source = json.loads((folder / "correction_input.json").read_text(encoding="utf8"))
                rows = [
                    {
                        "id": item["id"],
                        "corrected_text": "" if item["id"].startswith("m0_") else "corrected",
                        "selected_context_ids": [],
                        "uncertain_phrases": [],
                        "sensitive_changes": [],
                    }
                    for item in source["items"]
                ]
                app.save(folder / "correction_batch.json", {"result": {"items": rows}})
                return 0

            with (
                patch.object(app, "run_correction", side_effect=request) as runner,
                patch.object(app, "export", return_value=folder / "chat_final.xlsx"),
            ):
                app.process(
                    folder,
                    {"model": "test", "chatgpt_python": "python", "chatgpt_script": "unused"},
                    skip_asr=True,
                )
            cache = json.loads((folder / "correction_results.json").read_text(encoding="utf8"))
            self.assertEqual(cache["m0_a0"]["status"], "pending")
            self.assertEqual(cache["m1_a0"]["corrected_text"], "corrected")
            self.assertEqual(runner.call_count, 2)

    def test_partial_response_preserved_and_block_not_retried(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            (folder / "downloads").mkdir()
            messages = []
            raw = {}
            for i in range(3):
                ref = f"downloads/{i}.ogg"
                (folder / ref).write_bytes(str(i).encode())
                messages.append({"id": str(i), "text": "", "attachment_results": [ref]})
                raw[ref] = {
                    "status": "done",
                    "sha256": app.digest(folder / ref),
                    "text": "raw voice",
                }
            app.save(folder / "messages.json", messages)
            app.save(folder / "transcription_results.json", raw)

            def request(command, target):
                items = json.loads((folder / "correction_input.json").read_text(encoding="utf8"))[
                    "items"
                ]
                rows = [
                    dict(
                        id=item["id"],
                        corrected_text="corrected",
                        selected_context_ids=[],
                        uncertain_phrases=[],
                        sensitive_changes=[],
                    )
                    for item in items[:2]
                ]
                app.save(
                    folder / "correction_batch.json",
                    {
                        "result": {"items": rows},
                        "pending": [
                            {"id": items[2]["id"], "error": "content_filter", "status": "blocked"}
                        ],
                    },
                )
                return 0

            config = {"model": "test", "chatgpt_python": "python", "chatgpt_script": "unused"}
            with (
                patch.object(app, "run_correction", side_effect=request) as runner,
                patch.object(app, "export", return_value=folder / "chat_final.xlsx"),
            ):
                app.process(folder, config, skip_asr=True)
                app.process(folder, config, skip_asr=True)
                self.assertEqual(runner.call_count, 1)
            cache = json.loads((folder / "correction_results.json").read_text(encoding="utf8"))
            self.assertEqual(
                [cache[f"m{i}_a0"]["status"] for i in range(3)], ["done", "done", "blocked"]
            )
            self.assertEqual(cache["m2_a0"]["error"], "content_filter")

    def test_changed_audio_not_reused(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            (folder / "downloads").mkdir()
            path = folder / "downloads/v.ogg"
            path.write_bytes(b"audio")
            messages = [{"id": "1", "text": "", "attachment_results": ["downloads/v.ogg"]}]
            stale = {"downloads/v.ogg": {"status": "done", "sha256": "wrong", "text": "old"}}
            self.assertEqual(app.tasks(messages, folder, stale), [])
            stale["downloads/v.ogg"]["sha256"] = app.digest(path)
            work = app.tasks(messages, folder, stale)
            self.assertEqual(len(work), 1)
            self.assertEqual(len(work[0]["item"]["alternatives"]), 1)
            with self.assertRaises(ValueError):
                app.attachment(folder, "downloads/../../escape.ogg")

    def test_documented_export_long_text_links_and_pending(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            (folder / "downloads").mkdir()
            audio = folder / "downloads/v.ogg"
            audio.write_bytes(b"audio")
            messages = [
                {
                    "id": "1",
                    "sender": "=bad()",
                    "date": "نامشخص",
                    "time": "09:00",
                    "text": "x" * 33000,
                    "media": ["voice"],
                    "_media_done": True,
                    "attachment_results": ["downloads/v.ogg"],
                },
                {
                    "id": "2",
                    "text": "",
                    "media": ["photo"],
                    "attachment_results": ["download failed"],
                },
            ]
            raw = {
                "downloads/v.ogg": {
                    "status": "done",
                    "sha256": app.digest(audio),
                    "text": "raw not final",
                }
            }
            work = app.tasks(messages, folder, raw)
            corrections = {
                work[0]["item"]["id"]: {
                    "status": "done",
                    "input_hash": work[0]["input_hash"],
                    "corrected_text": "final voice",
                    "uncertain_phrases": ["uncertain"],
                }
            }
            target = app.export(messages, folder, raw, corrections, work)
            book = load_workbook(target)
            chat, files = book["Messages"], book["Files"]
            parts = (len(messages[0]["text"] + "\n\nfinal voice") + 899) // 900
            self.assertEqual(chat.max_row, 7 + parts)
            content = "".join(chat.cell(row, 6).value or "" for row in range(7, 7 + parts))
            self.assertIn("final voice", content)
            self.assertNotIn("raw not final", content)
            self.assertEqual(chat["D7"].data_type, "s")
            self.assertEqual(chat["B7"].value, "نامشخص")
            self.assertIn("اصلاح ماشینی", chat["H7"].value)
            self.assertIn("دانلود ناموفق", chat.cell(7 + parts, 9).value)
            self.assertEqual(files["G7"].hyperlink.target, audio.as_uri())
            self.assertEqual(chat["G7"].hyperlink.target, audio.as_uri())
            self.assertEqual(chat["G7"].value, "پخش ویس")
            self.assertEqual(files["A7"].hyperlink.target, "#'Messages'!A7")
            self.assertEqual(files["A8"].hyperlink.target, f"#'Messages'!A{7 + parts}")
            self.assertEqual(files["I8"].value, "دانلود ناموفق: download failed")
            self.assertTrue(chat.sheet_view.rightToLeft)
            self.assertEqual(chat.freeze_panes, "E7")
            self.assertIn("uncertain", chat["I7"].value)
            self.assertEqual(chat["J7"].value, "1")
            self.assertEqual(chat["K8"].value, 2)
            self.assertEqual(chat.auto_filter.ref, f"A6:K{7 + parts}")
            self.assertEqual(chat["E8"].value, "ادامهٔ پیام")
            self.assertFalse(chat.sheet_view.showGridLines)

    def test_unprocessed_voice_sensitive_guess_and_source_order(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            (folder / "downloads").mkdir()
            path = folder / "downloads/v.ogg"
            path.write_bytes(b"audio")
            messages = [
                {
                    "id": "late",
                    "time": "23:10",
                    "date": "نامشخص",
                    "sender": "من",
                    "media": [{"kind": "voice"}],
                    "attachment_results": ["downloads/v.ogg"],
                },
                {"id": "early", "time": "01:00", "sender": "Ali", "text": "next day?"},
            ]
            target = app.export(messages, folder, {}, {}, [])
            chat = load_workbook(target)["Messages"]
            self.assertEqual(chat["J7"].value, "late")
            self.assertEqual(chat["J8"].value, "early")
            self.assertEqual(chat["E7"].value, "ویس / صوت")
            self.assertIn("هنوز آماده نشده", chat["F7"].value)
            self.assertIn("آماده نیست", chat["H7"].value)
            raw = {"downloads/v.ogg": {"status": "done", "sha256": app.digest(path), "text": "raw"}}
            work = app.tasks(messages, folder, raw)
            corrections = {
                work[0]["item"]["id"]: {
                    "status": "done",
                    "input_hash": work[0]["input_hash"],
                    "corrected_text": "guess",
                    "sensitive_changes": [{"before": "name1", "after": "name2"}],
                }
            }
            target = app.export(messages, folder, raw, corrections, work)
            chat = load_workbook(target)["Messages"]
            self.assertIn("تغییر حساس", chat["I7"].value)
            self.assertIn("name2", chat["I7"].value)


if __name__ == "__main__":
    unittest.main()
