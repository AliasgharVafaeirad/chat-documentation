"""Readable evidence-oriented Excel view for the existing backup exporter."""

from audio_policy import is_voice, is_audio_file
import math
import re
from copy import copy
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

HEADER = 6
AUDIO = {".ogg", ".opus", ".mp3", ".wav", ".m4a", ".aac", ".flac", ".wma"}


def safe(value):
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def missing_date(message):
    return (
        not message.get("date")
        or message["date"] == "نامشخص"
        or message.get("date_source") == "telegram_day_header_partial"
    )


def resolved_sender(message):
    result = dict(message)
    if result.get("sender_role") == "self" and (
        not result.get("sender") or result.get("sender") == "نامشخص"
    ):
        result.update(sender="من", sender_source="outgoing_direction")
    return result


def role_sender(message):
    return (
        message.get("sender_status") == "unknown"
        or not message.get("sender")
        or message["sender"] in ("من", "طرف مقابل", "مخاطب", "نامشخص")
    )


def message_kind(message, refs):
    kinds = {m.get("kind", "") for m in message.get("media", []) if isinstance(m, dict)}
    if any(is_voice(message, ref) for ref in refs) or kinds == {"voice"}:
        return "ویس / صوت"
    labels = {
        "photo": "عکس",
        "image": "عکس",
        "video": "ویدیو",
        "document": "فایل",
        "sticker": "استیکر",
        "audio": "صوت",
        "animation": "انیمیشن",
    }
    return " + ".join(dict.fromkeys(labels.get(kind, kind) for kind in sorted(kinds) if kind)) or (
        "پیوست" if refs else "متن"
    )


def layout(sheet, title, notes, headers, widths, freeze):
    count = len(headers)
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    # Only presentation notes span empty cells; the actual table never uses merges.
    for row, text in [(2, title)] + [(3 + index, text) for index, text in enumerate(notes)]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=min(9, count))
        cell = sheet.cell(row, 1, text)
        cell.font = Font(
            name="Arial",
            size=15 if row == 2 else 11,
            bold=row == 2,
            color="20344A" if row == 2 else "526274",
        )
        cell.alignment = Alignment(
            horizontal="right", vertical="center", readingOrder=2, wrap_text=True
        )
        sheet.row_dimensions[row].height = 30
    for col, label in enumerate(headers, 1):
        cell = sheet.cell(HEADER, col, label)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476B")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True, readingOrder=2
        )
        cell.border = Border(right=Side(style="thin", color="FFFFFF"))
        sheet.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    sheet.row_dimensions[HEADER].height = 32
    sheet.freeze_panes = freeze
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:6"
    sheet.oddFooter.center.text = "Page &P / &N"


def finish(sheet, widths, last_column, centered):
    sheet.auto_filter.ref = f"A{HEADER}:{last_column}{max(HEADER, sheet.max_row)}"
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f"A1:{last_column}{sheet.max_row}"
    for row in sheet.iter_rows(min_row=HEADER + 1):
        lines = 1
        for cell in row:
            cell.font = Font(
                name="Arial",
                size=11,
                color="24476B" if cell.hyperlink else "20344A",
                underline="single" if cell.hyperlink else None,
            )
            cell.alignment = Alignment(
                horizontal="center" if cell.column in centered else "right",
                vertical="top",
                wrap_text=True,
                readingOrder=2,
            )
            stripe = int(row[0].value) if str(row[0].value or "").isdigit() else row[0].row - HEADER
            if stripe % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F6FA")
            # Estimate with each column's own width; narrow flags must not clip.
            available = max(4, widths[cell.column - 1] * 0.8)
            lines = max(
                lines,
                sum(
                    max(1, math.ceil(len(line) / available))
                    for line in str(cell.value or "").split("\n")
                ),
            )
        sheet.row_dimensions[row[0].row].height = min(409, max(38, lines * 18 + 8))


def write(
    messages, folder, whisper, corrections, work, attachment, save, output_name="chat_final.xlsx"
):
    messages = [resolved_sender(message) for message in messages]
    book = Workbook()
    chat = book.active
    chat.title = "Messages"
    files = book.create_sheet("Files")
    chat_widths = [8, 16, 12, 20, 14, 60, 23, 24, 35, 20, 8]
    file_widths = [8, 10, 16, 12, 20, 12, 38, 82, 32, 20, 16]
    unknown_dates = sum(missing_date(m) for m in messages)
    unknown_senders = sum(role_sender(m) for m in messages)
    voice_count = sum(
        any(is_voice(m, ref) for ref in m.get("attachment_results", []))
        or any(isinstance(v, dict) and v.get("kind") == "voice" for v in m.get("media", []))
        for m in messages
    )
    corrected_count = sum(
        corrections.get(t["item"]["id"], {}).get("status") == "done"
        and corrections[t["item"]["id"]].get("input_hash") == t["input_hash"]
        for t in work
    )
    layout(
        chat,
        "مستند گفتگو",
        [
            f"{len(messages)} پیام | {voice_count} پیام صوتی | {corrected_count} ویس آماده | پیام‌ها به ترتیب استخراج از چت آمده‌اند.",
            f"اطلاعات ناقص: تاریخ {unknown_dates} پیام مشخص نیست؛ برای {unknown_senders} پیام نام فرستنده ثبت نشده است.",
            "متن ویس با هوش مصنوعی تهیه شده؛ موارد مشکوک مشخص‌اند. برای مقایسه با اصل، روی «پخش ویس» کلیک کنید.",
        ],
        [
            "شماره پیام",
            "تاریخ",
            "ساعت",
            "فرستنده",
            "نوع پیام",
            "متن پیام",
            "فایل همراه",
            "وضعیت",
            "موارد نیازمند بررسی",
            "شناسه اصلی پیام",
            "بخش متن",
        ],
        chat_widths,
        "E7",
    )
    layout(
        files,
        "فایل‌های مرتبط با گفتگو",
        [
            "برای باز کردن فایل روی نام آن کلیک کنید. «بازگشت به پیام» شما را به پیام مربوط در گفتگو می‌برد.",
            "مسیر کامل مربوط به همین پوشه خروجی است؛ فایل ناموجود یا دانلود ناموفق، صریحاً مشخص شده است. حجم بر حسب بایت است.",
        ],
        [
            "پیام مربوط",
            "شماره فایل",
            "تاریخ",
            "ساعت",
            "فرستنده",
            "نوع فایل",
            "باز کردن فایل",
            "مسیر کامل فایل",
            "وضعیت فایل",
            "شناسه اصلی پیام",
            "حجم (بایت)",
        ],
        file_widths,
        "C7",
    )
    chat["D6"].comment = Comment(
        "«من» و «طرف مقابل» برچسب ثبت‌شده رابط هستند، نه نام احراز‌شده شخص. نام از داده منبع حدس زده نمی‌شود.",
        "User",
    )
    chat["B6"].comment = Comment(
        "تاریخ عین داده منبع حفظ می‌شود. تاریخ نامشخص از ساعت یا متن پیام استنباط نمی‌شود. در خروجی رسمی چت، زمان محلی منبع ثبت شده و منطقهٔ زمانی تأیید نشده است.",
        "User",
    )
    chat["A6"].comment = Comment(
        "ردیف ترتیب ثبت در messages.json است؛ بدون تاریخ کامل، ترتیب زمانی مستقل تأیید نشده است.",
        "User",
    )
    chat["J6"].comment = Comment(
        "برای تطبیق با messages.json. در خروجی واتساپ، این مقدار شمارهٔ رکورد فایل خروجی است؛ واتساپ در فایل متنی شناسهٔ اصلی پیام ارائه نمی‌کند.",
        "User",
    )
    chat["K6"].comment = Comment(
        "متن خیلی بلند در ردیف‌های متوالی با ردیف گفتگو و شناسه یکسان ادامه می‌یابد؛ پیام تازه نیست.",
        "User",
    )
    by_message = {}
    for task in work:
        by_message.setdefault(task["message_index"], []).append(task)
    enriched = []
    for index, message in enumerate(messages):
        refs = message.get("attachment_results", [])
        statuses, review, uncertainty, paths, names, voice_content = [], [], [], [], [], []
        first_file_row = files.max_row + 1
        voice_tasks = by_message.get(index, [])
        completed = 0
        for task in voice_tasks:
            result = corrections.get(task["item"]["id"], {})
            if result.get("input_hash") == task["input_hash"] and result.get("status") == "done":
                completed += 1
                voice_content.append(result["corrected_text"])
                uncertainty.extend(str(p) for p in result.get("uncertain_phrases", []))
                # Sensitive guesses must remain reviewable even if the model didn't mark uncertainty.
                for change in result.get("sensitive_changes", []):
                    if isinstance(change, dict):
                        review.append(
                            "تغییر حساس: "
                            + str(change.get("before", ""))
                            + " ← "
                            + str(change.get("after", ""))
                        )
            else:
                reason = (
                    result.get("error", "")
                    if result.get("input_hash") == task["input_hash"]
                    else ""
                )
                label = (
                    "توقف توسط فیلتر محتوایی سرویس؛ نیازمند بررسی"
                    if "content_filter" in reason
                    else "اصلاح آماده نیست"
                )
                statuses.append(
                    "ویس: "
                    + label
                    + (" | " + reason if reason and "content_filter" not in reason else "")
                )
                raw = task["item"]["alternatives"][0]["text"]
                voice_content.append("[متن خام Whisper — اصلاح نشده]\n" + raw)
        audio_count = 0
        for number, ref in enumerate(refs, 1):
            path = attachment(folder, ref)
            exists = path is not None and path.is_file()
            if path is not None:
                paths.append(str(path))
                names.append(path.name)
                audio_count += is_voice(message, ref)
                file_status = "ذخیره شده" if exists else "فایل موجود نیست"
                if not exists:
                    review.append("فایل موجود نیست: " + path.name)
                if is_audio_file(message, ref):
                    statuses.append("فایل صوتی: فقط بایگانی؛ تبدیل ویس انجام نمی‌شود")
                if is_voice(message, ref) and not any(t["ref"] == ref for t in voice_tasks):
                    statuses.append("ویس: تبدیل آماده نیست")
            else:
                names.append("دانلود ناموفق")
                file_status = "دانلود ناموفق: " + str(ref)
                review.append("دانلود ناموفق")
            files.append(
                [
                    "بازگشت به پیام",
                    number,
                    safe(message.get("date") or "نامشخص"),
                    safe(message.get("time") or "نامشخص"),
                    safe(message.get("sender") or "نامشخص"),
                    path.suffix.lower() if path else "نامشخص",
                    safe(path.name if path else ref),
                    safe(path or ""),
                    safe(file_status),
                    safe(message.get("id", "")),
                    path.stat().st_size if exists else None,
                ]
            )
            if exists:
                files.cell(files.max_row, 7).hyperlink = path.as_uri()
                files.cell(files.max_row, 8).hyperlink = path.as_uri()
            files.cell(files.max_row, 1).hyperlink = f"#'Messages'!A{chat.max_row + 1}"
        if message.get("media") and not refs:
            review.append("پیوست دانلود نشده")
        elif message.get("media") and message.get("_media_done") is False:
            review.append("دانلود پیوست کامل نشده")
        kind = message_kind(message, refs)
        if kind == "ویس / صوت" and not audio_count:
            statuses.append("ویس: فایل در دسترس نیست")
        if completed:
            statuses.insert(0, f"ویس: {completed}/{audio_count} تبدیل و اصلاح ماشینی")
        if not statuses:
            statuses.append(
                "متن ثبت شده" if not message.get("media") and not refs else "پیوست ثبت شده"
            )
        if uncertainty:
            review.append("عبارت نامطمئن: " + "؛ ".join(dict.fromkeys(uncertainty)))
        source_text = message.get("text", "").strip()
        if not message.get("sender_source"):
            review.append("نام فرستنده نیاز به استخراج مجدد دارد")
        elif role_sender(message):
            review.append("نام فرستنده قابل تشخیص نبود")
        if "edited" in str(message.get("time", "")).lower():
            review.append("ساعت دارای علامت ویرایش است")
        if not source_text and not refs and not message.get("media"):
            review.append("پیام بدون متن ثبت شده")
        content = "\n\n".join(([source_text] if source_text else []) + voice_content)
        if message.get("forwarded_from"):
            content = "فوروارد از: " + str(message["forwarded_from"]) + "\n\n" + content
        display_content = content or (
            "[متن ویس هنوز آماده نشده]"
            if kind == "ویس / صوت"
            else "[بدون متن؛ پیوست را ببینید]"
            if refs or message.get("media")
            else "[پیام بدون متن ثبت شده]"
        )
        single_file = attachment(folder, refs[0]) if len(refs) == 1 else None
        direct_file = single_file if single_file is not None and single_file.is_file() else None
        attachment_text = (
            ("پخش ویس" if direct_file.suffix.lower() in AUDIO else "باز کردن فایل")
            if direct_file
            else (f"مشاهده {len(refs)} فایل" if refs else "—")
        )
        values = [
            display_content,
            attachment_text,
            "؛ ".join(dict.fromkeys(statuses)),
            "؛ ".join(dict.fromkeys(review)) or "—",
        ]
        # Keep every continuation readable within Excel's 409-point row limit.
        limits = [900, 900, 650, 500]
        parts = max(1, max(math.ceil(len(value) / limit) for value, limit in zip(values, limits)))
        for part in range(parts):
            chunks = [
                safe(value[part * limit : (part + 1) * limit])
                for value, limit in zip(values, limits)
            ]
            chat.append(
                [
                    index + 1,
                    safe(message.get("date") or "نامشخص"),
                    safe(message.get("time") or "نامشخص"),
                    safe(message.get("sender") or "نامشخص"),
                    kind if part == 0 else "ادامهٔ پیام",
                    chunks[0],
                    chunks[1],
                    chunks[2],
                    chunks[3],
                    safe(message.get("id", "")),
                    part + 1,
                ]
            )
            if refs and part == 0:
                cell = chat.cell(chat.max_row, 7)
                cell.hyperlink = (
                    direct_file.as_uri() if direct_file else f"#'Files'!A{first_file_row}"
                )
                cell.comment = Comment("\n".join(names)[:30000], "User")
            if "edited" in str(message.get("time", "")).lower():
                chat.cell(chat.max_row, 3).comment = Comment(
                    "علامت edited در داده منبع وجود دارد؛ ساعت اصلی از ساعت ویرایش تفکیک نشده است.",
                    "User",
                )
        enriched.append(
            dict(
                message,
                final_content=content,
                processing_status=statuses,
                documentation_review=review,
                uncertain_phrases=uncertainty,
                full_file_paths=paths,
            )
        )
    finish(chat, chat_widths, "K", {1, 2, 3, 5, 11})
    finish(files, file_widths, "K", {1, 2, 3, 4, 6, 11})
    # Internal references are retained for tracing, but aren't reading columns.
    for column in ("A", "J", "K"):
        chat.column_dimensions[column].hidden = True
    files.column_dimensions["A"].width = 19
    files.column_dimensions["B"].hidden = True
    files.column_dimensions["J"].hidden = True
    chat.print_area = f"B1:I{chat.max_row}"
    for number in (2, 3, 4, 5):
        note = chat.cell(number, 1).value
        chat.unmerge_cells(f"A{number}:I{number}")
        chat.merge_cells(f"B{number}:I{number}")
        chat.cell(number, 2, note)
        chat.cell(number, 2).font = copy(chat.cell(number, 1).font)
        chat.cell(number, 2).alignment = copy(chat.cell(number, 1).alignment)
        chat.cell(number, 1).value = None
    for row in chat.iter_rows(min_row=7):
        message = messages[row[0].value - 1]
        source = message.get("sender_source", "legacy_unknown")
        row[3].comment = Comment(
            "منبع نام فرستنده: "
            + source
            + "\nشناسه فرستنده: "
            + str(message.get("sender_id") or "ثبت نشده")
            + "\nجهت پیام: "
            + str(message.get("sender_role") or "ثبت نشده"),
            "User",
        )
    for row in files.iter_rows(min_row=7):
        for col in (7, 8, 10):
            row[col - 1].alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True, readingOrder=1
            )
    for row in chat.iter_rows(min_row=7):
        row[2].alignment = Alignment(
            horizontal="center", vertical="top", wrap_text=True, readingOrder=1
        )
    amber = PatternFill("solid", fgColor="FFF1CC")
    red = PatternFill("solid", fgColor="FCE4E4")
    if messages:
        end = chat.max_row
        chat.conditional_formatting.add(
            f"B7:B{end}", FormulaRule(formula=['B7="نامشخص"'], fill=amber)
        )
        chat.conditional_formatting.add(
            f"D7:D{end}",
            FormulaRule(formula=['OR(D7="من",D7="طرف مقابل",D7="نامشخص")'], fill=amber),
        )
        chat.conditional_formatting.add(
            f"I7:I{end}", FormulaRule(formula=['AND(I7<>"—",I7<>"")'], fill=amber)
        )
        chat.conditional_formatting.add(
            f"H7:H{end}", FormulaRule(formula=['ISNUMBER(SEARCH("آماده نیست",H7))'], fill=amber)
        )
    if files.max_row > HEADER:
        files.conditional_formatting.add(
            f"I7:I{files.max_row}", FormulaRule(formula=['I7<>"ذخیره شده"'], fill=red)
        )
        files.conditional_formatting.add(
            f"C7:C{files.max_row}", FormulaRule(formula=['C7="نامشخص"'], fill=amber)
        )
        files.conditional_formatting.add(
            f"E7:E{files.max_row}",
            FormulaRule(formula=['OR(E7="من",E7="طرف مقابل",E7="نامشخص")'], fill=amber),
        )
    files.column_dimensions["K"].number_format = "#,##0"
    if Path(output_name).name != output_name or not output_name.endswith(".xlsx"):
        raise ValueError("Invalid Excel output name")
    target = folder / output_name
    temporary = target.with_suffix(".tmp.xlsx")
    book.save(temporary)
    try:
        temporary.replace(target)
    except PermissionError:
        target = folder / ("chat_final_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f") + ".xlsx")
        temporary.replace(target)
        print("Existing Excel is open or locked. New export saved as: " + str(target), flush=True)
    save(folder / "messages_final.json", enriched)
    if folder.parent.name.endswith("_backups"):
        from report_library import publish_excel

        try:
            target = publish_excel(folder, target)
        except PermissionError:
            print(
                "Readable Excel name is open; current export saved at: " + str(target), flush=True
            )
    return target
